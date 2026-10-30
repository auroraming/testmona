"""Reversible XLSX exchange for the versioned test-case mind-map format."""

from __future__ import annotations

from io import BytesIO
from typing import Any
from zipfile import BadZipFile, ZipFile

from fastapi import Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from .. import auth
from ..database import get_db
from ..models import User
from .base import router
from .helpers import ensure_import_permission, get_test_suite_or_404

MAX_XLSX_BYTES = 10 * 1024 * 1024
MAX_XLSX_ROWS = 20_000
MAX_XLSX_EXPANDED_BYTES = 50 * 1024 * 1024
FOLDER_HEADERS = ("folder_key", "parent_folder_key", "name")
CASE_HEADERS = (
    "case_key", "folder_key", "title", "priority", "test_type", "description", "preconditions", "is_multistep",
)
STEP_HEADERS = ("case_key", "step_number", "action", "expected_result", "step_type")


def _sheet_rows(sheet: Worksheet, expected_headers: tuple[str, ...]) -> list[dict[str, Any]]:
    iterator = sheet.iter_rows(values_only=False)
    try:
        header_cells = next(iterator)
    except StopIteration:
        raise HTTPException(status_code=400, detail=f"Worksheet '{sheet.title}' is empty")
    headers = tuple(str(cell.value or "").strip() for cell in header_cells)
    if headers != expected_headers:
        raise HTTPException(
            status_code=400,
            detail=f"Worksheet '{sheet.title}' headers must be: {', '.join(expected_headers)}",
        )
    rows: list[dict[str, Any]] = []
    for row_number, cells in enumerate(iterator, start=2):
        if row_number > MAX_XLSX_ROWS + 1:
            raise HTTPException(status_code=400, detail=f"XLSX is limited to {MAX_XLSX_ROWS} rows per worksheet")
        if any(cell.data_type == "f" for cell in cells):
            raise HTTPException(status_code=400, detail=f"Worksheet '{sheet.title}' row {row_number} contains a formula")
        values = [cell.value for cell in cells[:len(expected_headers)]]
        if not any(value not in (None, "") for value in values):
            continue
        rows.append(dict(zip(expected_headers, values)))
    return rows


def workbook_to_mindmap_document(content: bytes) -> tuple[dict[str, Any], list[str]]:
    if len(content) > MAX_XLSX_BYTES:
        raise HTTPException(status_code=413, detail="XLSX files are limited to 10 MB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 10_000 or sum(member.file_size for member in members) > MAX_XLSX_EXPANDED_BYTES:
                raise HTTPException(status_code=413, detail="Expanded XLSX content is limited to 50 MB")
            if any(member.filename.startswith(("/", "\\")) or ".." in member.filename.split("/") for member in members):
                raise HTTPException(status_code=400, detail="XLSX contains an unsafe archive path")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except HTTPException:
        raise
    except (BadZipFile, KeyError, ValueError, OSError) as exc:
        raise HTTPException(status_code=400, detail="Invalid XLSX workbook") from exc
    try:
        required = {"Folders", "Test Cases", "Steps"}
        if not required.issubset(workbook.sheetnames):
            raise HTTPException(status_code=400, detail="XLSX requires Folders, Test Cases, and Steps worksheets")
        folder_rows = _sheet_rows(workbook["Folders"], FOLDER_HEADERS)
        case_rows = _sheet_rows(workbook["Test Cases"], CASE_HEADERS)
        step_rows = _sheet_rows(workbook["Steps"], STEP_HEADERS)
    finally:
        workbook.close()
    warnings: list[str] = []
    folder_by_key: dict[str, dict[str, Any]] = {}
    for index, row in enumerate(folder_rows, start=1):
        key = str(row["folder_key"] or "").strip()
        name = str(row["name"] or "").strip()
        if not key or key in folder_by_key or not name:
            raise HTTPException(status_code=400, detail=f"Folders row {index + 1} requires a unique folder_key and name")
        folder_by_key[key] = {"id": index, "name": name, "parent_key": str(row["parent_folder_key"] or "").strip() or None, "children": []}
    roots: list[dict[str, Any]] = []
    for key, folder in folder_by_key.items():
        parent_key = folder.pop("parent_key")
        if parent_key:
            parent = folder_by_key.get(parent_key)
            if not parent:
                raise HTTPException(status_code=400, detail=f"Folder '{key}' references unknown parent '{parent_key}'")
            parent["children"].append(folder)
        else:
            roots.append(folder)

    def visit(folder: dict[str, Any], ancestors: set[int]) -> None:
        if folder["id"] in ancestors:
            raise HTTPException(status_code=400, detail="Folder hierarchy contains a cycle")
        for child in folder["children"]:
            visit(child, {*ancestors, folder["id"]})

    for root in roots:
        visit(root, set())
    if len({folder["id"] for root in roots for folder in _walk(root)}) != len(folder_by_key):
        raise HTTPException(status_code=400, detail="Folder hierarchy contains a cycle")

    steps_by_case: dict[str, list[dict[str, Any]]] = {}
    for row_number, row in enumerate(step_rows, start=2):
        case_key = str(row["case_key"] or "").strip()
        if not case_key:
            raise HTTPException(status_code=400, detail=f"Steps row {row_number} requires case_key")
        steps_by_case.setdefault(case_key, []).append({
            "step_number": row["step_number"] or len(steps_by_case.get(case_key, [])) + 1,
            "action": str(row["action"] or ""),
            "expected_result": str(row["expected_result"] or ""),
            "step_type": str(row["step_type"] or "manual"),
        })

    cases: list[dict[str, Any]] = []
    seen_case_keys: set[str] = set()
    for row_number, row in enumerate(case_rows, start=2):
        case_key = str(row["case_key"] or "").strip()
        title = str(row["title"] or "").strip()
        if not case_key or case_key in seen_case_keys or not title:
            raise HTTPException(status_code=400, detail=f"Test Cases row {row_number} requires a unique case_key and title")
        seen_case_keys.add(case_key)
        folder_key = str(row["folder_key"] or "").strip()
        if folder_key and folder_key not in folder_by_key:
            warnings.append(f"{title}: folder_key '{folder_key}' was not found; case will be imported at suite root")
        value = str(row["is_multistep"] or "").strip().lower()
        is_multistep = value in {"1", "true", "yes", "y"}
        steps = steps_by_case.pop(case_key, [])
        cases.append({
            "id": case_key,
            "section_id": folder_by_key[folder_key]["id"] if folder_key in folder_by_key else None,
            "title": title,
            "priority": str(row["priority"] or "medium"),
            "test_type": str(row["test_type"] or "manual"),
            "description": str(row["description"] or ""),
            "preconditions": str(row["preconditions"] or ""),
            "is_multistep": is_multistep,
            "steps": steps,
        })
    if steps_by_case:
        raise HTTPException(status_code=400, detail=f"Steps reference unknown case_key: {next(iter(steps_by_case))}")
    return {
        "format": "testmona.test-case-mindmap",
        "version": 1,
        "sections": roots,
        "test_cases": cases,
    }, warnings


def _walk(folder: dict[str, Any]):
    yield folder
    for child in folder["children"]:
        yield from _walk(child)


def _template_bytes() -> bytes:
    workbook = Workbook()
    folders = workbook.active
    folders.title = "Folders"
    folders.append(FOLDER_HEADERS)
    folders.append(("checkout", "", "Checkout"))
    folders.append(("refund", "checkout", "Refund"))
    cases = workbook.create_sheet("Test Cases")
    cases.append(CASE_HEADERS)
    cases.append(("TC-1", "refund", "Refund a paid order", "high", "manual", "Verify refund", "Order is paid", True))
    steps = workbook.create_sheet("Steps")
    steps.append(STEP_HEADERS)
    steps.append(("TC-1", 1, "Open the order", "Order details open", "manual"))
    steps.append(("TC-1", 2, "Submit refund", "Refund is created", "verification"))
    instructions = workbook.create_sheet("Instructions")
    instructions.append(("TestMona reversible XLSX template",))
    instructions.append(("Keep worksheet names and header rows unchanged. Keys are file-local identifiers and are never trusted as database IDs.",))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@router.get("/templates/test-case-mindmap.xlsx")
def download_mindmap_xlsx_template(current_user: User = Depends(auth.get_current_active_user)):
    del current_user
    return Response(
        content=_template_bytes(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="testmona-test-case-template.xlsx"'},
    )


@router.post("/import/test-cases/mindmap-xlsx/convert")
async def convert_mindmap_xlsx(
    test_suite_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(auth.get_current_active_user),
):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Only .xlsx files are supported")
    suite = get_test_suite_or_404(db, test_suite_id)
    ensure_import_permission(current_user, suite.project_id, db)
    document, warnings = workbook_to_mindmap_document(await file.read(MAX_XLSX_BYTES + 1))
    return {"document": document, "warnings": warnings}
